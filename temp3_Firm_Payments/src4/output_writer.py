
"""Writes results into a copy of the template."""
from copy import copy
from pathlib import Path

import openpyxl

from temp3_Firm_Payments.src4.settings import DATE_FMT, FIRST_DATA_ROW, MONEY_FMT


def write_output(template: Path, output: Path, results: list[dict]) -> None:
    wb = openpyxl.load_workbook(template)
    ws = wb.worksheets[0]

    if ws.max_row >= FIRST_DATA_ROW:  # clear hint rows under the header
        ws.delete_rows(FIRST_DATA_ROW, ws.max_row - FIRST_DATA_ROW + 1)

    formats = [DATE_FMT, "General", "General", MONEY_FMT, MONEY_FMT, MONEY_FMT, "General"]  # same 2-dp display as the original; stored values keep 3 dp
    for row_no, r in enumerate(results, start=FIRST_DATA_ROW):
        values = [r["date"], r["account"], r["name"],
                  float(r["gross"]), float(r["net"]), float(r["commission"]), r["origination"]]
        for col, (value, fmt) in enumerate(zip(values, formats), start=1):
            c = ws.cell(row=row_no, column=col, value=value)
            c.number_format = fmt
            header = ws.cell(row=1, column=col)
            font = copy(header.font)
            font.name = "Arial Narrow"
            c.font = font
            c.alignment = copy(header.alignment)

    output.parent.mkdir(parents=True, exist_ok=True)
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                font = copy(cell.font)
                font.name = "Arial Narrow"
                cell.font = font
    wb.save(output)