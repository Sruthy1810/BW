from pathlib import Path
from datetime import datetime
import shutil
from copy import copy
import re

import pandas as pd
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import PatternFill
from notify.fonts import apply_font_to_file
from notify.mail import record_files
from config.config_loader import load_temp3_Court_config
from logs.logger_setup import setup_logger
from temp1.src1.t1_parser import format_name
from temp3_Court_Cost.src3.court_cost_parser import (
    JUDGMENT_HIGHLIGHT_COLUMN,
    OUTPUT_COLUMNS,
    read_costs,
)

JUDGMENT_ROW_FILL = PatternFill(fill_type="solid", fgColor="FFFFFF00")


def _move_input_file(input_file, destination_folder, log, status):
    source = Path(input_file)
    destination_dir = Path(destination_folder)
    destination_dir.mkdir(parents=True, exist_ok=True)

    destination = destination_dir / source.name
    suffix = 1
    while destination.exists() and destination.resolve() != source.resolve():
        destination = destination_dir / f"{source.stem} ({suffix}){source.suffix}"
        suffix += 1

    if source.resolve() != destination.resolve():
        shutil.move(str(source), str(destination))
    log.info("%s input %s to %s", status, source, destination)
    return destination


def _header_key(value):
    return re.sub(r"[^a-z0-9]", "", str(value).strip().lower())


def _find_output_header(workbook):
    required_columns = [column for column in OUTPUT_COLUMNS if column != "Judgment Date"]
    expected = {_header_key(column): column for column in required_columns}
    for worksheet in workbook.worksheets:
        for row_number in range(1, worksheet.max_row + 1):
            found = {
                _header_key(cell.value): cell.column
                for cell in worksheet[row_number]
                if cell.value is not None
            }
            if expected.keys() <= found.keys():
                output_columns = {
                    column: found[_header_key(column)] for column in required_columns
                }
                judgment_key = _header_key("Judgment Date")
                if judgment_key in found:
                    output_columns["Judgment Date"] = found[judgment_key]
                else:
                    insert_at = output_columns["Placement date"] + 1
                    worksheet.insert_cols(insert_at)
                    output_columns = {
                        column: (
                            column_index + 1
                            if column_index >= insert_at else column_index
                        )
                        for column, column_index in output_columns.items()
                    }
                    output_columns["Judgment Date"] = insert_at
                    worksheet.cell(
                        row=row_number,
                        column=insert_at,
                        value="Judgment Date",
                    )
                    source_header = worksheet.cell(
                        row=row_number,
                        column=insert_at - 1,
                    )
                    target_header = worksheet.cell(
                        row=row_number,
                        column=insert_at,
                    )
                    if source_header.has_style:
                        target_header._style = copy(source_header._style)
                    target_header.number_format = source_header.number_format
                    target_header.alignment = copy(source_header.alignment)
                    target_header.font = copy(source_header.font)
                    target_header.fill = copy(source_header.fill)
                    target_header.border = copy(source_header.border)
                return worksheet, row_number, output_columns
    raise ValueError(
        "Output template is missing a header row with: "
        + ", ".join(required_columns)
    )


def _write_output(data, template_file, output_file, log):
    template = Path(template_file)
    output = Path(output_file)
    if not template.is_file():
        raise FileNotFoundError(f"Court cost output template not found: {template}")
    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(template, output)

    workbook = load_workbook(output)
    try:
        worksheet, header_row, output_columns = _find_output_header(workbook)
        first_data_row = header_row + 1
        last_template_row = worksheet.max_row
        style_source = {
            column: copy(worksheet.cell(first_data_row, column)._style)
            for column in output_columns.values()
            if first_data_row <= last_template_row
        }
        clear_through = max(last_template_row, first_data_row + len(data) - 1)
        for row_number in range(first_data_row, clear_through + 1):
            for cell in worksheet[row_number]:
                if not isinstance(cell, MergedCell):
                    cell.value = None

        for row_offset, values in enumerate(
            data[OUTPUT_COLUMNS].itertuples(index=False, name=None)
        ):
            row_number = first_data_row + row_offset
            for column in output_columns.values():
                if row_number > last_template_row and column in style_source:
                    worksheet.cell(row_number, column)._style = copy(style_source[column])

            for column_name, value in zip(OUTPUT_COLUMNS, values):
                cell = worksheet.cell(row_number, output_columns[column_name])
                if column_name == "Debtor Name" and value is not None:
                    value = format_name(str(value))
                cell.value = value
                if column_name == "Account Number":
                    account_text = str(value).strip()
                    if re.fullmatch(r"\d+", account_text) and len(account_text) <= 15:
                        cell.value = int(account_text)
                        cell.number_format = "0000000000"
                if column_name in ("Effective Date", "Placement date", "Judgment Date") and value is not None:
                    cell.number_format = "dd-mm-yyyy"
                elif column_name == "Transaction Amt.":
                    cell.number_format = "0.00"

            if bool(data.iloc[row_offset].get(JUDGMENT_HIGHLIGHT_COLUMN, False)):
                for column in output_columns.values():
                    worksheet.cell(row_number, column).fill = copy(JUDGMENT_ROW_FILL)

        workbook.save(output)
    finally:
        workbook.close()
    log.info("Court cost output written: %d rows -> %s", len(data), output)


def _unique_output_path(output_folder, input_file):
    date_text = datetime.now().strftime("%m-%d-%Y")
    base = f"{input_file.stem.upper()} - OUTPUT - {date_text}"
    output = output_folder / f"{base}.xlsx"
    suffix = 1
    while output.exists():
        output = output_folder / f"{base} ({suffix}).xlsx"
        suffix += 1
    return output


def run_temp3_Court():
    config = load_temp3_Court_config()
    log = setup_logger(config["log_file"], config["log_level"])
    log.info("===== Temp3 Court Cost run started =====")

    input_folder = Path(config["input_folder"])
    files = sorted(
        path for path in input_folder.iterdir()
        if path.is_file()
        and path.suffix.lower() in (".xls", ".xlsx")
        and not path.name.startswith("~$")
    )
    done, failed = [], []
    log.info("Court cost input files found: %d", len(files))

    for input_file in files:
        log.info("----- Processing %s -----", input_file.name)
        try:
            data = read_costs(input_file, log)
            output_file = _unique_output_path(Path(config["output_folder"]), input_file)
            _write_output(data, config["output_template"], output_file, log)
            apply_font_to_file(output_file)          # <-- add here, LAST step on the output
            log.info("Arial Narrow applied to %s", output_file)

            _move_input_file(input_file, config["processed_folder"], log, "Processed")
            done.append(input_file.name)
        except Exception:
            log.exception("Failed to process court cost input: %s", input_file)
            _move_input_file(input_file, config["failed_folder"], log, "Failed")
            failed.append(input_file.name)

    record_files("temp3_court_cost", len(files), len(done), len(failed))
    log.info("Court cost files done: %d | failed: %d %s", len(done), len(failed), failed or "")
    log.info("===== Temp3 Court Cost run finished =====")
    if failed:
        raise RuntimeError(f"{len(failed)} court cost file(s) failed: {failed}")
