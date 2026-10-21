import logging
import re
from decimal import Decimal, InvalidOperation

import pandas as pd
from openpyxl import load_workbook

log = logging.getLogger("remittance")


def keep_valid_rows(df: pd.DataFrame, col_a: int, col_b: int,
                    rule_mode: str = "BOTH", fees_col: int = 7,
                    include_cost_only: bool = False,
                    fees_as_principal: bool = True,
                    skip_interest_only: bool = False,
                    include_other_only: bool = False) -> pd.DataFrame:
    a = df.iloc[:, col_a]
    b = df.iloc[:, col_b]
    mode = rule_mode.upper()

    if mode == "SUM":                 # Principal + Interest not 0
        mask = (a + b).round(2) != 0
        if fees_as_principal:         # Fees-only payment (current behaviour)
            fees = df.iloc[:, fees_col]
            mask = mask | ((a == 0) & (b == 0) & (fees != 0))
        if include_cost_only:         # only when enabled for this client
            mask = mask | (df.iloc[:, 8] != 0)
        if include_other_only:
            mask = mask | (df["Other"] != 0)
    elif mode == "ANY":               # either one is non-zero
        mask = (a != 0) | (b != 0)
    else:                             # BOTH: both columns non-zero
        mask = (a != 0) & (b != 0)

    if skip_interest_only:            # interest-only row: not a principal payment
        fees = df.iloc[:, fees_col]
        costs = df.iloc[:, 8]
        interest_only = (a == 0) & (fees == 0) & (costs == 0)
        mask = mask & ~interest_only

    valid = df[mask].reset_index(drop=True)
    log.info("Rule: %s | columns %d and %d | total: %d | valid: %d | removed: %d",
             mode, col_a, col_b, len(df), len(valid), len(df) - len(valid))
    return valid


def remove_empty_data_sheets(path, log=None):
    """Remove empty/all-zero Temp2 data sheets, preserving Summary sheets."""
    wb = load_workbook(path)
    sheet_specs = {
        "payments": (9, (4, 5, 6, 7, 8, 9)),
        "cost": (8, (3,)),
        "intonly": (8, (3,)),
    }
    removed = []

    for ws in list(wb.worksheets):
        if "summary" in ws.title.lower():
            continue
        spec = sheet_specs.get(ws.title.strip().lower())
        if spec is None:
            continue

        data_start, amount_columns = spec
        data_rows = [
            row for row in ws.iter_rows(min_row=data_start, max_col=ws.max_column)
            if any(cell.value is not None and str(cell.value).strip() for cell in row)
        ]
        if data_rows:
            has_nonzero = False
            for row in data_rows:
                for column in amount_columns:
                    value = row[column - 1].value if column <= len(row) else None
                    if value is None or (isinstance(value, str) and not value.strip()):
                        continue
                    if isinstance(value, str) and value.startswith("="):
                        has_nonzero = True
                        break
                    try:
                        amount = Decimal(str(value).replace("$", "").replace(",", "").strip())
                    except InvalidOperation:
                        has_nonzero = True
                        break
                    if amount != 0:
                        has_nonzero = True
                        break
                if has_nonzero:
                    break
            if has_nonzero:
                continue

        removed.append(ws.title)
        wb.remove(ws)

    for sheet_name in removed:
        reference = re.compile(
            rf"(?<![A-Za-z0-9_])(?:'{re.escape(sheet_name)}'|{re.escape(sheet_name)})!",
            re.IGNORECASE,
        )
        for ws in wb.worksheets:
            if "summary" not in ws.title.lower():
                continue
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("=") and reference.search(cell.value):
                        cell.value = "=0"
                        if log:
                            log.info("Summary formula %s!%s set to zero after removing '%s'",
                                     ws.title, cell.coordinate, sheet_name)

    if removed:
        wb.save(path)
    wb.close()
    if log and removed:
        log.info("Removed empty/all-zero Temp2 sheets: %s", removed)