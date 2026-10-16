from copy import copy
from pathlib import Path
import shutil
from datetime import datetime
import pandas as pd
from openpyxl import load_workbook


PAYMENT_COLUMNS = 10
COST_COLUMNS = 6
INT_ONLY_COLUMNS = 3


def write_output(output_file: str, template_file: str, payments: pd.DataFrame,
                 costs: pd.DataFrame, statement_summary: pd.DataFrame,
                 client_share: float,input_file:str, bw_share: float,
                 bw_interest_share: float, remittance_address: str) -> None:
    output_path = Path(output_file)
    template_path = Path(template_file)
    if not template_path.is_file():
        raise FileNotFoundError(f"Output template was not found: {template_path}")
    if output_path.resolve() == template_path.resolve():
        raise ValueError("The output file cannot be the template file")

    payment_rows = _payment_rows(
        payments, client_share, bw_share, bw_interest_share
    )
    cost_rows = _cost_rows(payments, costs)
    interest_only_rows = _interest_only_rows(payments, bw_interest_share)
    expected_sheets = {"Payments", "Cost", "IntOnly"}

    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template_path, output_path)
    workbook = load_workbook(output_path)
    missing_sheets = expected_sheets.difference(workbook.sheetnames)
    if missing_sheets:
        raise ValueError(
            f"Template is missing required worksheet(s): {', '.join(sorted(missing_sheets))}"
        )

    payments_sheet = workbook["Payments"]
    payments_sheet["G4"] = client_share        # above CLIENT SHARE
    payments_sheet["H4"] = bw_share            # above BW, LLP SHARE
    payments_sheet["G4"].number_format = "0.0%"
    payments_sheet["H4"].number_format = "0.0%"
    payments_sheet["A1"] = input_file.stem.upper()
    payments_sheet["A2"]= "Payments Through " + datetime.now().strftime("%m-%d-%Y")
    



    cost_sheets= workbook["Cost"]
    cost_sheets["B1"]= input_file.stem.upper()
    cost_sheets["B2"]= "Costs Through " + datetime.now().strftime("%m-%d-%Y")

    intOnly_Sheet = workbook["IntOnly"]
    intOnly_Sheet["A1"]=input_file.stem.upper()
    intOnly_Sheet["A2"]= "Interest Through " + datetime.now().strftime("%m-%d-%Y")

    address_lines = [line.strip() for line in remittance_address.splitlines() if line.strip()]
    if len(address_lines) > 3:
        address_lines = address_lines[:2] + [" ".join(address_lines[2:])]
    address_lines.extend([""] * (3 - len(address_lines)))

    for sheet_name in ("Summary1", "Summary2", "Summary3", "Summary4"):
        summary_sheet = workbook[sheet_name]
        summary_sheet["A15"] = address_lines[0]
        summary_sheet["A16"] = address_lines[1]
        summary_sheet["A17"] = address_lines[2]
        summary_sheet["A20"] = "Statements Through " + datetime.now().strftime("%m-%d-%Y")



    _write_rows(payments_sheet, payment_rows, start_row=9, style_row=9,
                column_count=PAYMENT_COLUMNS)
    _write_rows(workbook["Cost"], cost_rows, start_row=8, style_row=8,
                column_count=COST_COLUMNS)
    _write_rows(workbook["IntOnly"], interest_only_rows, start_row=8, style_row=8,
                column_count=INT_ONLY_COLUMNS)

    for sheet_name, rows in (
        ("Payments", payment_rows),
        ("Cost", cost_rows),
        ("IntOnly", interest_only_rows),
    ):
        if rows:
            continue
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    formula = cell.value
                    if (
                        isinstance(formula, str)
                        and formula.startswith("=")
                        and (
                            f"{sheet_name}!" in formula
                            or f"'{sheet_name}'!" in formula
                        )
                    ):
                        cell.value = "=0"
        workbook.remove(workbook[sheet_name])

    if "Summary1" in workbook.sheetnames:
        workbook["Summary1"]["D27"] = -_direct_payment_commission(statement_summary)

    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.save(output_path)


def _payment_rows(payments: pd.DataFrame, client_share: float, bw_share: float,
                  bw_interest_share: float) -> list[tuple]:
    interest_only = (payments["Principal"] == 0) & (payments["Interest"] != 0)
    applied_principal = (
        payments["Principal"] + payments["Fees"] - payments["Direct Payments"]
    )
    total_payment = (
        applied_principal
        + payments["Interest"]
        + payments["Costs"]
        + payments["Other"]
    )
    has_payment = total_payment != 0
    selected = payments.loc[has_payment & ~interest_only].copy()
    selected["Applied Principal"] = applied_principal.loc[selected.index]
    selected["Total Payment"] = total_payment.loc[selected.index]
    selected["Display Name"] = selected["Debtor Name"].map(_format_defendant_name)
    selected = selected.sort_values("Display Name", kind="stable")

    return [
        (
            row["Display Name"],
            row["Plaintiff"],
            row["Ref Acct #"],
            row["Total Payment"],
            row["Applied Principal"],
            row["Costs"],
            row["Applied Principal"] * client_share,        # G  CLIENT SHARE
            row["Applied Principal"] * bw_share,            # H  BW, LLP SHARE
            row["Interest"] * bw_interest_share,            # I  BW, LLP INTEREST (unchanged)
            "",
        )
        for _, row in selected.iterrows()
    ]


def _cost_rows(payments: pd.DataFrame, costs: pd.DataFrame) -> list[tuple]:
    plaintiff_by_claim = payments.drop_duplicates("Claim").set_index("Claim")["Plaintiff"]
    rows = []
    for _, row in costs.iterrows():
        rows.append(
            (
                _format_defendant_name(row["Debtor Name"]),
                plaintiff_by_claim.get(row["Claim"], ""),
                row["Account #"],
                row["Amount"],
                row["Cost Description"],
                row["Notes"] or None,
            )
        )
    return sorted(rows, key=lambda row: row[0].casefold())


def _interest_only_rows(payments: pd.DataFrame,
                        bw_interest_share: float) -> list[tuple]:
    selected = payments.loc[
        (payments["Principal"] == 0) & (payments["Interest"] != 0)
    ].copy()
    selected["Display Name"] = selected["Debtor Name"].map(_format_defendant_name)
    selected = selected.sort_values("Display Name", kind="stable")
    return [
        (
            row["Display Name"],
            row["Plaintiff"],
            row["Interest"] * bw_interest_share,
        )
        for _, row in selected.iterrows()
    ]


def _write_rows(worksheet, rows: list[tuple], start_row: int, style_row: int,
                column_count: int) -> None:
    existing_last_row = max(worksheet.max_row, start_row + len(rows) - 1)
    for row_number in range(start_row, existing_last_row + 1):
        for column in range(1, column_count + 1):
            cell = worksheet.cell(row=row_number, column=column)
            cell.value = None

    for offset, values in enumerate(rows):
        row_number = start_row + offset
        for column in range(1, column_count + 1):
            cell = worksheet.cell(row=row_number, column=column)
            if row_number != style_row:
                source = worksheet.cell(row=style_row, column=column)
                if source.has_style:
                    cell._style = copy(source._style)
                if source.number_format:
                    cell.number_format = source.number_format
            cell.value = values[column - 1]


def _direct_payment_commission(statement_summary: pd.DataFrame) -> float:
    if statement_summary.empty:
        return 0.0
    matching = statement_summary.loc[
        statement_summary["Metric"].str.strip().str.casefold()
        == "fees due for direct payments"
    ]
    if matching.empty:
        return 0.0
    return float(matching.iloc[-1]["Amount"])


def _format_defendant_name(name: str) -> str:
    parts = str(name).strip().split()
    if len(parts) < 2:
        return str(name).strip()
    suffixes = {"jr.", "sr.", "ii", "iii", "iv"}
    surname_end = len(parts) - 1
    if parts[-1].lower() in suffixes and len(parts) > 2:
        surname_end -= 1
    surname = " ".join(parts[surname_end:])
    given_names = " ".join(parts[:surname_end])
    return f"{surname}, {given_names}"
