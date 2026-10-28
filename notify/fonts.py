from copy import copy
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

FONT_NAME = "Arial Narrow"


def _with_font(font, name=FONT_NAME):
    f = copy(font)
    f.name = name
    f.scheme = None        # without this, Excel keeps the theme font (Calibri)
    return f


def apply_font_to_wb(wb, name=FONT_NAME):
    # default style (best effort; some versions don't allow it)
    try:
        for style in wb._named_styles:
            style.font = _with_font(style.font, name)
    except Exception:
        pass

    for ws in wb.worksheets:
        # every existing cell
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row,
                                min_col=1, max_col=ws.max_column):
            for cell in row:
                if isinstance(cell, MergedCell):
                    continue
                cell.font = _with_font(cell.font, name)

        # column / row defaults, so empty cells and new typing use it too
        for dim in ws.column_dimensions.values():
            dim.font = _with_font(dim.font, name)
        for dim in ws.row_dimensions.values():
            dim.font = _with_font(dim.font, name)


def apply_font_to_file(path):
    wb = load_workbook(path)
    apply_font_to_wb(wb)
    wb.save(path)
    wb.close()