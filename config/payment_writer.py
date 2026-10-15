import logging
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

log = logging.getLogger("remittance")
MONEY_FMT = '"$"#,##0.00'
HEADERS = ["DEFENDANT'S NAME", "PLAINTIFF", "REF NO.", "TOTAL PAYMENT",
           "APPLIED TO PRINCIPAL", "COSTS RETURNED", "CLIENT SHARE",
           "BW, LLP SHARE", "BW, LLP INTEREST", "Notes"]


def format_payment_sheet(path, df, cfg):
    thru = cfg["pay_thru_date"]
    if thru.upper() == "AUTO":
        thru = datetime.now().strftime("%m-%d-%Y")
    pct = cfg["client_share_pct"]

    rows = []
    for _, r in df.iterrows():
        
        principal = round(float(r.iloc[cfg["pay_principal_col"]]), 2)
        client = round(principal * pct, 2)
        if (principal+client ==0):
            continue
        rows.append([
            r.iloc[1], r.iloc[13], r.iloc[2],
            round(float(r.iloc[cfg["pay_total_col"]]), 2),
            principal,
            round(float(r.iloc[cfg["pay_costs_col"]]), 2),
            client,
            round(principal - client, 2),
            round(float(r.iloc[cfg["pay_interest_col"]]), 2),
            "",
        ])

    wb = load_workbook(path)
    old = wb.worksheets[0]
    name = old.title
    wb.remove(old)
    ws = wb.create_sheet(name, 0)
    bold = Font(bold=True)

    ws["A1"] = cfg["pay_title"].upper()
    ws["A1"].font = bold
    ws["A2"] = f"Payments Thru {thru}"
    ws["A2"].font = bold

    ws["G4"] = pct
    ws["H4"] = round(1 - pct, 4)
    for c in ("G4", "H4"):
        ws[c].number_format = "0%"
        ws[c].font = bold

    ws["C5"] = "Total"
    ws["C5"].font = bold
    for col in range(4, 10):                      # D..I
        total = round(sum(r[col - 1] for r in rows), 2)
        cell = ws.cell(row=5, column=col, value=total)
        cell.number_format = MONEY_FMT
        cell.font = bold

    for c, h in enumerate(HEADERS, start=1):
        ws.cell(row=7, column=c, value=h).font = bold

    for i, row in enumerate(rows, start=8):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=i, column=c, value=v)
            if 4 <= c <= 9:
                cell.number_format = MONEY_FMT

    for i, w in enumerate([26, 45, 12, 16, 22, 16, 15, 15, 18, 18], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    wb.save(path)
    log.info("First sheet formatted with %d rows", len(rows))