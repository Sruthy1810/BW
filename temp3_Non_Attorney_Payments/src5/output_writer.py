"""Writes results into a copy of the template.
The column layout is detected from the template's header row (see settings.LAYOUTS)."""
import logging
from copy import copy
from pathlib import Path

import openpyxl

from temp3_Non_Attorney_Payments.src5.settings import FIRST_DATA_ROW, LAYOUTS

log = logging.getLogger("remittance")


def _norm(v) -> str:
    return " ".join(str(v or "").split()).lower()


def detect_layout(ws) -> tuple[str, list]:
    found = [_norm(c.value) for c in ws[1] if c.value is not None]
    for name, layout in LAYOUTS.items():
        if found == [_norm(h) for h in layout["headers"]]:
            return name, layout["columns"]
    raise ValueError(f"Template headers not recognised: {found}. Add a layout in settings.LAYOUTS.")


def write_output(template: Path, output: Path, results: list[dict]) -> None:
    wb = openpyxl.load_workbook(template)
    ws = wb.worksheets[0]
    layout_name, columns = detect_layout(ws)
    log.info("Template layout detected: %s", layout_name)

    if ws.max_row >= FIRST_DATA_ROW:  # clear hint rows under the header
        ws.delete_rows(FIRST_DATA_ROW, ws.max_row - FIRST_DATA_ROW + 1)

    for row_no, r in enumerate(results, start=FIRST_DATA_ROW):
        for col, (field, fmt) in enumerate(columns, start=1):
            value = r[field]
            if field in ("gross", "net", "commission"):
                value = float(value)
            c = ws.cell(row=row_no, column=col, value=value)
            c.number_format = fmt
            header = ws.cell(row=1, column=col)
            font = copy(header.font)
            font.name = "Arial Narrow"
            c.font = font
            c.alignment = copy(header.alignment)

    output.parent.mkdir(parents=True, exist_ok=True)
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                font = copy(cell.font)
                font.name = "Arial Narrow"
                cell.font = font
    wb.save(output)