"""Reads the first sheet of the raw report (.xls or .xlsx) into a list of rows."""
from pathlib import Path

import openpyxl


def read_rows(path: Path) -> list[list]:
    ext = path.suffix.lower()
    if ext == ".xls":
        return _read_xls(path)
    if ext in (".xlsx", ".xlsm"):
        ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    raise ValueError(f"Unsupported input file type: {ext}")


def _read_xls(path: Path) -> list[list]:
    try:
        import xlrd
    except ImportError:
        raise SystemExit("Reading .xls needs xlrd:  pip install xlrd")
    wb = xlrd.open_workbook(path)
    sh = wb.sheet_by_index(0)
    rows = []
    for r in range(sh.nrows):
        row = []
        for c in range(sh.ncols):
            cell = sh.cell(r, c)
            if cell.ctype == xlrd.XL_CELL_DATE:
                row.append(xlrd.xldate_as_datetime(cell.value, wb.datemode))
            elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                row.append(None)
            else:
                row.append(cell.value)
        rows.append(row)
    return rows